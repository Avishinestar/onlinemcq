import openpyxl
import json
import csv
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('Employability_Skills_Marathi_English_580_Questions.xlsx')
ws = wb.active
print(f"Excel Sheet Title: {ws.title}")
print(f"Excel Total Rows (Header + Data): {ws.max_row}")
print(f"Excel Total Columns: {ws.max_column}")

with open('Employability_Skills_Marathi_English_580_Questions.json', encoding='utf-8') as f:
    data = json.load(f)
print(f"JSON Total Questions: {len(data)}")

print("\n--- SAMPLE PREVIEWS ---")
for q in data[:3]:
    print(f"Q{q['sr_no']} [{q['module_name_mr']} | {q['lesson_name_mr']}]:")
    print(f"  EN Q: {q['question_en']}")
    print(f"  MR Q: {q['question_mr']}")
    print(f"  Options:")
    print(f"    A: {q['option_a_en']}  -->  {q['option_a_mr']}")
    print(f"    B: {q['option_b_en']}  -->  {q['option_b_mr']}")
    print(f"    C: {q['option_c_en']}  -->  {q['option_c_mr']}")
    print(f"    D: {q['option_d_en']}  -->  {q['option_d_mr']}")
    print(f"  Correct Answer: [{q['correct_answer_letter']}] {q['correct_answer_en']}  -->  {q['correct_answer_mr']}")
    print("-" * 50)

# Check question on page 41 Neena
q14_candidates = [q for q in data if 'Neena' in q['question_en']]
if q14_candidates:
    q14 = q14_candidates[0]
    print(f"Special Check (Page 41 Question):")
    print(f"  Q{q14['sr_no']} Q.No.{q14['q_no']}: {q14['question_en']}")
    print(f"  MR: {q14['question_mr']}")
    print(f"  Ans: [{q14['correct_answer_letter']}] {q14['correct_answer_mr']}")
    print("-" * 50)

# Verify no empty fields
empty_fields = 0
for idx, q in enumerate(data, 1):
    for k, v in q.items():
        if v is None or str(v).strip() == '':
            print(f"Empty field {k} in item {idx}")
            empty_fields += 1
print(f"Total Empty Fields in Entire Dataset: {empty_fields}")
