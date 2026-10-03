import pdfplumber
import re
import json
import time
import os
import urllib.request
import urllib.parse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

CACHE_FILE = "translation_cache.json"

def clean(t):
    if not t: return ''
    return re.sub(r'\s+', ' ', str(t)).strip()

def load_cache():
    if os.path.exists(CACHE_FILE):
        try:
            with open(CACHE_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}

def save_cache(cache):
    with open(CACHE_FILE, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=2)

def translate_single(text, retries=3):
    text_clean = text.strip()
    if not text_clean:
        return ""
    if re.match(r'^[0-9\s\.,\-\+/\(\)]+$', text_clean):
        return text_clean
        
    url = 'https://translate.googleapis.com/translate_a/single?client=gtx&sl=en&tl=mr&dt=t&q=' + urllib.parse.quote(text_clean)
    for attempt in range(retries):
        try:
            req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'})
            with urllib.request.urlopen(req, timeout=10) as resp:
                data = json.loads(resp.read().decode('utf-8'))
                return ''.join(item[0] for item in data[0] if item[0]).strip()
        except Exception as e:
            time.sleep(1 + attempt * 2)
    return text_clean

def translate_batch(texts, cache):
    to_fetch = [t for t in texts if t and t not in cache and not re.match(r'^[0-9\s\.,\-\+/\(\)]+$', t)]
    if not to_fetch:
        return
        
    delimiter = '\n=====\n'
    batch_size = 20
    
    for i in range(0, len(to_fetch), batch_size):
        chunk = to_fetch[i:i+batch_size]
        combined = delimiter.join(chunk)
        url = 'https://translate.googleapis.com/translate_a/single?client=gtx&sl=en&tl=mr&dt=t&q=' + urllib.parse.quote(combined)
        
        success = False
        try:
            req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'})
            with urllib.request.urlopen(req, timeout=15) as resp:
                data = json.loads(resp.read().decode('utf-8'))
                full_res = ''.join(item[0] for item in data[0] if item[0])
                parts = [s.strip() for s in full_res.split('=====')]
                if len(parts) == len(chunk):
                    for orig, trans in zip(chunk, parts):
                        cache[orig] = trans
                    success = True
                else:
                    print(f"Batch count mismatch ({len(parts)} vs {len(chunk)}). Falling back to individual translation.")
        except Exception as e:
            print(f"Batch request error: {e}. Falling back to individual.")
            
        if not success:
            for item in chunk:
                cache[item] = translate_single(item)
                time.sleep(0.3)
                
        save_cache(cache)
        print(f"Translated {min(i+batch_size, len(to_fetch))}/{len(to_fetch)} new unique strings...")
        time.sleep(0.8)

# Domain-specific glossary refinement for ITI trainees
GLOSSARY_REFINEMENTS = {
    "All of these": "वरील सर्व",
    "None of these": "यापैकी काहीही नाही",
    "Employability Skills": "रोजगार कौशल्ये (Employability Skills)",
    "Introduction to Employability Skills": "रोजगार कौशल्यांची ओळख (Introduction to Employability Skills)",
    "Constitutional Values - Citizenship": "घटनात्मक मूल्ये - नागरिकत्व",
    "Constitutional Values-Citizenship": "घटनात्मक मूल्ये - नागरिकत्व",
    "Becoming a Professional in the 21st Century": "२१ व्या शतकात व्यावसायिक बनणे",
    "Basic English Skills": "मूलभूत इंग्रजी कौशल्ये",
    "Communication Skills": "संभाषण कौशल्ये (Communication Skills)",
    "Essential Digital Skills": "आवश्यक डिजिटल कौशल्ये",
    "Financial and Legal Literacy": "आर्थिक आणि कायदेशीर साक्षरता",
    "Diversity and Inclusion": "विविधता आणि समावेशन",
    "Entrepreneurship": "उद्योजकता (Entrepreneurship)",
    "Career Development and Goal Setting": "करिअर विकास आणि ध्येय निश्चिती",
    "Customer Service": "ग्राहक सेवा (Customer Service)",
    "Getting Ready for Apprenticeships & Jobs": "प्रशिक्षणार्थी (Apprenticeship) आणि नोकरीसाठी पूर्वतयारी",
    "Getting Ready for Apprenticeship & Jobs": "प्रशिक्षणार्थी (Apprenticeship) आणि नोकरीसाठी पूर्वतयारी"
}

def refine_translation(orig_en, trans_mr):
    if orig_en in GLOSSARY_REFINEMENTS:
        return GLOSSARY_REFINEMENTS[orig_en]
    
    # Common fixes
    res = trans_mr
    res = res.replace("सर्व या", "वरील सर्व")
    res = res.replace("या सर्व", "वरील सर्व")
    return res

def extract_all_questions():
    pdf_path = "Employability Skills 2022 - 1st year.pdf"
    questions = []
    current_module = ''
    current_lesson = ''
    
    with pdfplumber.open(pdf_path) as pdf:
        for p_no in range(3, len(pdf.pages) + 1):
            page = pdf.pages[p_no - 1]
            text = page.extract_text() or ''
            mod_m = re.search(r'Module Name\s*:\s*(.*)', text)
            if mod_m:
                raw_m = mod_m.group(1).split('Correct')[0].strip()
                current_module = raw_m
                
            tables = page.extract_tables()
            for t in tables:
                for r in t:
                    start_i = 0
                    while start_i < len(r) and r[start_i] is None: start_i += 1
                    end_i = len(r)
                    while end_i > start_i and r[end_i-1] is None: end_i -= 1
                    row = r[start_i:end_i]
                    if not row: continue
                    if any('Option A' in str(c) or 'Correct Answer' in str(c) for c in row):
                        continue
                    ans = clean(row[-1]).upper() if row else ''
                    if ans not in ['A', 'B', 'C', 'D']:
                        continue
                    
                    lesson = current_lesson
                    q_no = ''
                    q_text = ''
                    optA = clean(row[-5])
                    optB = clean(row[-4])
                    optC = clean(row[-3])
                    optD = clean(row[-2])
                    
                    lead = row[:-5]
                    lead_clean = [clean(c) for c in lead]
                    if len(lead_clean) == 3:
                        if lead_clean[0]: lesson = lead_clean[0]; current_lesson = lesson
                        q_no = lead_clean[1]
                        q_text = lead_clean[2]
                    elif len(lead_clean) == 2:
                        if re.match(r'^\d+$', lead_clean[0]):
                            q_no = lead_clean[0]
                            q_text = lead_clean[1]
                        elif re.match(r'^(\d+)\s+(.*)', lead_clean[0]):
                            m = re.match(r'^(\d+)\s+(.*)', lead_clean[0])
                            q_no = m.group(1)
                            q_text = m.group(2) + (' ' + lead_clean[1] if lead_clean[1] else '')
                        elif re.match(r'^(\d+)\s+(.*)', lead_clean[1]):
                            if lead_clean[0]: lesson = lead_clean[0]; current_lesson = lesson
                            m = re.match(r'^(\d+)\s+(.*)', lead_clean[1])
                            q_no = m.group(1)
                            q_text = m.group(2)
                    elif len(lead_clean) == 1:
                        m = re.match(r'^(\d+)\s+(.*)', lead_clean[0])
                        if m:
                            q_no = m.group(1)
                            q_text = m.group(2)
                            
                    # Get correct answer text
                    ans_map = {'A': optA, 'B': optB, 'C': optC, 'D': optD}
                    ans_text = ans_map.get(ans, '')
                    
                    questions.append({
                        'sr_no': len(questions) + 1,
                        'page': p_no,
                        'module': current_module,
                        'lesson': lesson,
                        'q_no': q_no,
                        'question': q_text,
                        'option_a': optA,
                        'option_b': optB,
                        'option_c': optC,
                        'option_d': optD,
                        'answer_letter': ans,
                        'answer_text': ans_text
                    })
    return questions

def main():
    print("Extracting 580 questions from PDF...")
    questions = extract_all_questions()
    print(f"Extracted {len(questions)} questions.")
    assert len(questions) == 580, f"Expected 580 questions, got {len(questions)}"
    
    # Collect all unique strings
    cache = load_cache()
    all_strings = set()
    for q in questions:
        all_strings.add(q['module'])
        all_strings.add(q['lesson'])
        all_strings.add(q['question'])
        all_strings.add(q['option_a'])
        all_strings.add(q['option_b'])
        all_strings.add(q['option_c'])
        all_strings.add(q['option_d'])
        all_strings.add(q['answer_text'])
    
    needed = [s for s in all_strings if s and s not in cache]
    print(f"Total unique strings: {len(all_strings)}, need translation: {len(needed)}")
    
    if needed:
        print("Starting batch translation to Marathi...")
        translate_batch(needed, cache)
    
    # Apply refinements & build complete dataset
    final_rows = []
    for q in questions:
        mod_en = q['module']
        mod_mr = refine_translation(mod_en, cache.get(mod_en, mod_en))
        
        les_en = q['lesson']
        les_mr = refine_translation(les_en, cache.get(les_en, les_en))
        
        q_en = q['question']
        q_mr = refine_translation(q_en, cache.get(q_en, q_en))
        
        oa_en = q['option_a']
        oa_mr = refine_translation(oa_en, cache.get(oa_en, oa_en))
        
        ob_en = q['option_b']
        ob_mr = refine_translation(ob_en, cache.get(ob_en, ob_en))
        
        oc_en = q['option_c']
        oc_mr = refine_translation(oc_en, cache.get(oc_en, oc_en))
        
        od_en = q['option_d']
        od_mr = refine_translation(od_en, cache.get(od_en, od_en))
        
        ans_letter = q['answer_letter']
        ans_map_mr = {'A': oa_mr, 'B': ob_mr, 'C': oc_mr, 'D': od_mr}
        ans_mr = ans_map_mr.get(ans_letter, cache.get(q['answer_text'], q['answer_text']))
        
        final_rows.append({
            'sr_no': q['sr_no'],
            'page_no': q['page'],
            'module_name_en': mod_en,
            'module_name_mr': mod_mr,
            'lesson_name_en': les_en,
            'lesson_name_mr': les_mr,
            'q_no': q['q_no'],
            'question_en': q_en,
            'question_mr': q_mr,
            'option_a_en': oa_en,
            'option_a_mr': oa_mr,
            'option_b_en': ob_en,
            'option_b_mr': ob_mr,
            'option_c_en': oc_en,
            'option_c_mr': oc_mr,
            'option_d_en': od_en,
            'option_d_mr': od_mr,
            'correct_answer_letter': ans_letter,
            'correct_answer_en': q['answer_text'],
            'correct_answer_mr': ans_mr
        })
        
    # Save JSON
    json_path = "Employability_Skills_Marathi_English_580_Questions.json"
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(final_rows, f, ensure_ascii=False, indent=2)
    print(f"Saved JSON to {json_path}")
    
    # Save CSV (UTF-8 with BOM for 100% correct display in Excel / Google Sheets)
    import csv
    csv_path = "Employability_Skills_Marathi_English_580_Questions.csv"
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(final_rows[0].keys()))
        writer.writeheader()
        writer.writerows(final_rows)
    print(f"Saved CSV to {csv_path}")
    
    # Save Formatted Excel (.xlsx)
    xlsx_path = "Employability_Skills_Marathi_English_580_Questions.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Employability Skills 580 MCQs"
    
    headers = [
        "Sr. No. / क्र.",
        "Module Name (English)",
        "Module Name (मराठी)",
        "Lesson Name (English)",
        "Lesson Name (मराठी)",
        "Q. No.",
        "Question (English)",
        "Question (मराठी)",
        "Option A (English)",
        "Option A (मराठी)",
        "Option B (English)",
        "Option B (मराठी)",
        "Option C (English)",
        "Option C (मराठी)",
        "Option D (English)",
        "Option D (मराठी)",
        "Correct Answer (A/B/C/D)",
        "Correct Answer (English)",
        "Correct Answer (मराठी)"
    ]
    
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid") # Dark Navy Blue
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
    ws.row_dimensions[1].height = 28
    
    alt_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    
    for r_idx, row in enumerate(final_rows, start=2):
        row_vals = [
            row['sr_no'],
            row['module_name_en'],
            row['module_name_mr'],
            row['lesson_name_en'],
            row['lesson_name_mr'],
            row['q_no'],
            row['question_en'],
            row['question_mr'],
            row['option_a_en'],
            row['option_a_mr'],
            row['option_b_en'],
            row['option_b_mr'],
            row['option_c_en'],
            row['option_c_mr'],
            row['option_d_en'],
            row['option_d_mr'],
            row['correct_answer_letter'],
            row['correct_answer_en'],
            row['correct_answer_mr']
        ]
        ws.append(row_vals)
        ws.row_dimensions[r_idx].height = 36
        
        is_even = (r_idx % 2 == 0)
        for c_idx in range(1, len(row_vals) + 1):
            c = ws.cell(row=r_idx, column=c_idx)
            c.border = thin_border
            if is_even:
                c.fill = alt_fill
            if c_idx in [1, 6, 17]:
                c.alignment = center_align
            else:
                c.alignment = left_align
                
    # Column width adjustment
    col_widths = {
        1: 8,   # Sr No
        2: 22,  # Module EN
        3: 22,  # Module MR
        4: 20,  # Lesson EN
        5: 20,  # Lesson MR
        6: 8,   # Q No
        7: 35,  # Question EN
        8: 35,  # Question MR
        9: 20,  # Opt A EN
        10: 20, # Opt A MR
        11: 20, # Opt B EN
        12: 20, # Opt B MR
        13: 20, # Opt C EN
        14: 20, # Opt C MR
        15: 20, # Opt D EN
        16: 20, # Opt D MR
        17: 12, # Correct Ans
        18: 25, # Correct Ans EN
        19: 25  # Correct Ans MR
    }
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width
        
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    
    wb.save(xlsx_path)
    print(f"Saved styled Excel workbook to {xlsx_path}")
    print("ALL DONE SUCCESSFULLY!")

if __name__ == "__main__":
    main()
